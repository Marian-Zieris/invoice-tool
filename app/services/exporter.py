import os
import uuid
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook

from app.models import Invoice

EXPORT_DIR = os.environ.get("EXPORT_DIR") or os.path.join(os.getcwd(), "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)

# Interní klíč -> výchozí hlavička sloupce, pokud zákazník nemá vlastní excel_template_config.
DEFAULT_COLUMNS: Dict[str, str] = {
    "invoice_id": "ID faktury",
    "original_filename": "Soubor",
    "supplier_name": "Dodavatel",
    "invoice_date": "Datum",
    "description": "Popis položky",
    "category": "Kategorie",
    "amount": "Částka",
    "currency": "Měna",
    "confidence_score": "Jistota",
    "is_corrected": "Opraveno ručně",
}


def _row_for_line_item(invoice: Invoice, item) -> Dict[str, Any]:
    return {
        "invoice_id": invoice.id,
        "original_filename": invoice.original_filename,
        "supplier_name": invoice.supplier_name or "",
        "invoice_date": invoice.invoice_date.isoformat() if invoice.invoice_date else "",
        "description": item.description,
        "category": item.category,
        "amount": item.amount,
        "currency": invoice.currency,
        "confidence_score": item.confidence_score,
        "is_corrected": item.is_corrected,
    }


def _build_dataframe(invoices: List[Invoice], columns: List[str]) -> pd.DataFrame:
    rows = []
    for invoice in invoices:
        for item in invoice.line_items:
            row = _row_for_line_item(invoice, item)
            rows.append({key: row.get(key, "") for key in columns})
    return pd.DataFrame(rows, columns=columns)


def export_invoices_to_excel(invoices: List[Invoice], excel_template_config: Optional[Dict[str, Any]] = None) -> str:
    """Vyexportuje LineItem záznamy vybraných faktur do .xlsx podle šablony zákazníka."""
    config = excel_template_config if isinstance(excel_template_config, dict) else {}

    columns: List[str] = config.get("columns") or list(DEFAULT_COLUMNS.keys())
    headers: List[str] = config.get("headers") or [DEFAULT_COLUMNS.get(col, col) for col in columns]

    dataframe = _build_dataframe(invoices, columns)

    output_name = f"export_{uuid.uuid4()}.xlsx"
    output_path = os.path.join(EXPORT_DIR, output_name)

    template_path = config.get("template_path")
    if template_path and os.path.exists(template_path):
        workbook = load_workbook(template_path)
        sheet_name = config.get("sheet_name") or workbook.sheetnames[0]
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
        start_row = int(config.get("start_row", 2))

        for col_idx, header in enumerate(headers, start=1):
            sheet.cell(row=start_row - 1, column=col_idx, value=header)

        for row_offset, row in enumerate(dataframe.itertuples(index=False)):
            for col_idx, value in enumerate(row, start=1):
                sheet.cell(row=start_row + row_offset, column=col_idx, value=value)

        workbook.save(output_path)
    else:
        dataframe.columns = headers
        dataframe.to_excel(output_path, index=False, engine="openpyxl")

    return output_path
